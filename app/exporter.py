"""Export reviewed records to CSV / Excel.

Excel has two sheets that mirror the planned Lark Base tables:
  - "Header":   one row per scanned form
  - "Chi tiet": one row per measurement, linked by record_id
"""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.templates_loader import Template

STATUS_VI = {"pass": "Đạt", "fail": "Không đạt", "invalid": "Không đọc được", "n/a": ""}
META_COLS = [("record_id", "Mã bản ghi"), ("form", "Mã form"), ("uploaded_at", "Thời gian upload"),
             ("uploaded_by", "Người upload"), ("reviewed_by", "Người duyệt"),
             ("reviewed_at", "Thời gian duyệt"), ("overall", "Kết luận")]


def _label(f) -> str:
    return f"{f.label} ({f.unit})" if f.unit else f.label


def _value(section: dict, key: str) -> str:
    return (section.get(key) or {}).get("value", "")


def header_rows(tpl: Template, docs: list[dict]) -> tuple[list[str], list[list]]:
    cols = [label for _, label in META_COLS] + [_label(f) for f in tpl.header_fields] + ["Ghi chú AI"]
    rows = []
    for d in docs:
        data = d.get("data") or {}
        header = data.get("header", {})
        rows.append(
            [d["id"], tpl.code, d["uploaded_at"], d["uploaded_by"], d.get("reviewed_by") or "",
             d.get("reviewed_at") or "", STATUS_VI.get(d.get("overall") or "n/a", "")]
            + [_value(header, f.key) for f in tpl.header_fields]
            + [data.get("notes", "")]
        )
    return cols, rows


def detail_rows(tpl: Template, docs: list[dict]) -> tuple[list[str], list[list]]:
    limited = [c for c in tpl.columns if c.min is not None or c.max is not None]
    cols = (["Mã bản ghi", "Số lô", "Dòng"] + [_label(c) for c in tpl.columns]
            + [f"Đánh giá {c.label}" for c in limited])
    rows = []
    for d in docs:
        data = d.get("data") or {}
        evaluation = d.get("evaluation") or {}
        batch_no = _value(data.get("header", {}), "batch_no")
        ev_rows = evaluation.get("rows", [])
        for i, row in enumerate(data.get("rows", [])):
            ev = ev_rows[i] if i < len(ev_rows) else {}
            rows.append(
                [d["id"], batch_no, i + 1]
                + [_value(row, c.key) for c in tpl.columns]
                + [STATUS_VI.get(ev.get(c.key, "n/a"), "") for c in limited]
            )
    return cols, rows


def to_csv(tpl: Template, docs: list[dict]) -> bytes:
    """Flat CSV: header values repeated on each measurement row (easy to import as one table)."""
    h_cols, h_rows = header_rows(tpl, docs)
    d_cols, d_rows = detail_rows(tpl, docs)
    by_id: dict[int, list[list]] = {}
    for r in d_rows:
        by_id.setdefault(r[0], []).append(r[2:])  # drop record_id, batch_no (already in header)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(h_cols + d_cols[2:])
    for h in h_rows:
        details = by_id.get(h[0]) or [[""] * (len(d_cols) - 2)]
        for det in details:
            w.writerow(h + det)
    return out.getvalue().encode("utf-8-sig")  # BOM để Excel hiển thị đúng tiếng Việt


def to_xlsx(tpl: Template, docs: list[dict]) -> bytes:
    wb = Workbook()
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2F5597")
    for idx, (title, (cols, rows)) in enumerate(
        [("Header", header_rows(tpl, docs)), ("Chi tiet", detail_rows(tpl, docs))]
    ):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = title
        ws.append(cols)
        for cell in ws[1]:
            cell.font, cell.fill = bold, fill
        for r in rows:
            ws.append(r)
        for i, col in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(40, len(col) + 4))
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
