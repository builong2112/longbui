import io

import openpyxl
import pymupdf

from app import db

from .conftest import SAMPLE_WC, cell, make_jpeg


def upload(client, files, template_code="IPC-WC"):
    r = client.post("/api/upload", data={"template_code": template_code}, files=files)
    assert r.status_code == 200, r.text
    return r.json()["document_ids"]


def test_requires_login(fake_claude):
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as c:
        assert c.get("/api/documents").status_code == 401


def test_upload_extract_review_approve_export(client):
    [doc_id] = upload(client, [("files", ("form.jpg", make_jpeg(), "image/jpeg"))])

    doc = client.get(f"/api/documents/{doc_id}").json()
    assert doc["status"] == "pending_review"
    assert doc["overall"] == "fail"
    assert doc["ai_data"] == doc["data"]

    # IPC corrects the misread weight
    data = doc["data"]
    data["rows"][1]["net_weight"] = cell("49,7", "verified")
    r = client.put(f"/api/documents/{doc_id}/data", json={"data": data})
    assert r.json()["changed"] == 1
    assert r.json()["evaluation"]["overall"] == "pass"

    doc = client.get(f"/api/documents/{doc_id}").json()
    assert doc["ai_data"]["rows"][1]["net_weight"]["value"] == "48,7"  # AI original kept
    [edit] = doc["history"]["edits"]
    assert (edit["field_path"], edit["old_value"], edit["new_value"], edit["edited_by"]) == (
        "rows.2.net_weight", "48,7", "49,7", "ipc1")

    assert client.post(f"/api/documents/{doc_id}/approve", json={}).status_code == 200
    assert client.get(f"/api/documents/{doc_id}").json()["status"] == "approved"

    # editing an approved record needs a reason and re-opens it
    data["header"]["batch_no"] = cell("L240999", "verified")
    assert client.put(f"/api/documents/{doc_id}/data", json={"data": data}).status_code == 400
    r = client.put(f"/api/documents/{doc_id}/data", json={"data": data, "reason": "Ghi sai số lô"})
    assert r.status_code == 200
    assert client.get(f"/api/documents/{doc_id}").json()["status"] == "pending_review"
    client.post(f"/api/documents/{doc_id}/approve", json={})

    r = client.get("/api/export", params={"template_code": "IPC-WC", "format": "xlsx"})
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Header", "Chi tiet"]
    header = list(wb["Header"].iter_rows(values_only=True))
    assert header[1][0] == doc_id and "L240999" in header[1]
    details = list(wb["Chi tiet"].iter_rows(values_only=True))
    assert len(details) == 3  # title + 2 measurement rows
    assert details[2][-1] == "Đạt"

    r = client.get("/api/export", params={"template_code": "IPC-WC", "format": "csv"})
    text = r.content.decode("utf-8-sig")
    assert text.count("\n") == 3 and "49,7" in text


def test_pdf_is_split_into_pages(client):
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.new_page()
    ids = upload(client, [("files", ("scan.pdf", pdf.tobytes(), "application/pdf"))])
    assert len(ids) == 2
    assert [client.get(f"/api/documents/{i}").json()["page_no"] for i in ids] == [1, 2]


def test_auto_detect_uses_classifier(client, fake_claude):
    [doc_id] = upload(client, [("files", ("a.jpg", make_jpeg(), "image/jpeg"))], template_code="auto")
    assert client.get(f"/api/documents/{doc_id}").json()["template_code"] == "IPC-WC"
    assert "form_code" in fake_claude.calls[0]["properties"]


def test_extraction_error_is_recorded_and_retry_works(client, fake_claude, monkeypatch):
    from app import extractor

    def boom(*args, **kwargs):
        raise extractor.ExtractionError("API key Anthropic không hợp lệ hoặc chưa cấu hình.")

    monkeypatch.setattr(extractor, "_call", boom)
    [doc_id] = upload(client, [("files", ("a.jpg", make_jpeg(), "image/jpeg"))])
    doc = client.get(f"/api/documents/{doc_id}").json()
    assert doc["status"] == "error" and "API key" in doc["error"]

    monkeypatch.setattr(extractor, "_call", fake_claude)
    assert client.post(f"/api/documents/{doc_id}/retry", json={}).status_code == 200
    assert client.get(f"/api/documents/{doc_id}").json()["status"] == "pending_review"


def test_reject_needs_reason(client):
    [doc_id] = upload(client, [("files", ("a.jpg", make_jpeg(), "image/jpeg"))])
    assert client.post(f"/api/documents/{doc_id}/reject", json={}).status_code == 400
    assert client.post(f"/api/documents/{doc_id}/reject", json={"note": "Ảnh mờ"}).status_code == 200
    assert client.get("/api/documents", params={"status": "rejected"}).json()[0]["id"] == doc_id


def test_unsupported_file_rejected(client):
    r = client.post("/api/upload", data={"template_code": "IPC-WC"},
                    files=[("files", ("a.docx", b"xx", "application/octet-stream"))])
    assert r.status_code == 400


def test_sanitize_drops_unknown_fields(client):
    [doc_id] = upload(client, [("files", ("a.jpg", make_jpeg(), "image/jpeg"))])
    data = {**SAMPLE_WC, "header": {**SAMPLE_WC["header"], "hacker": cell("x")}}
    client.put(f"/api/documents/{doc_id}/data", json={"data": data})
    assert "hacker" not in db.get_document(doc_id)["data"]["header"]
